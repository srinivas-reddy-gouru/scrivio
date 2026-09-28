"""Completion checks at the download boundary (review item R03).

The packaging button is disabled while placeholders remain, but a button
is a suggestion: the download URL answered anyone who asked. These tests
make the requests the UI would never make.
"""
import copy
import io

import docx
import pytest
from fastapi.testclient import TestClient

from api import server
from pipeline.schemas.models import ResumeDoc, StructuredResume, TailoredResume
from pipeline.workers.resume_studio_worker import render_markdown, run_ats_checks

ORIGINAL = {
    "basics": {"name": "Sam Okafor", "label": "Software Engineer",
               "email": "sam@example.com", "phone": "+1 555 010 9999",
               "summary": "Software engineer who builds internal platforms."},
    "work": [{"name": "Initrode", "position": "Software Engineer",
              "startDate": "Feb 2020", "endDate": "Present", "summary": "",
              "highlights": ["Maintained 12 services behind the internal gateway",
                             "Consolidated build clusters to reduce costs"]}],
    "skills": [{"name": "Languages", "keywords": ["Python"]}],
}
FORMATS = ("md", "json", "docx", "pdf")


@pytest.fixture
def desk(tmp_path, monkeypatch):
    monkeypatch.setattr(server, "OUTPUT_ROOT", tmp_path)
    original = StructuredResume.model_validate(ORIGINAL)
    unfinished = copy.deepcopy(ORIGINAL)
    unfinished["work"][0]["highlights"][1] = (
        "Consolidated build clusters, reducing costs by [METRIC]")
    doc = ResumeDoc(
        resume_id="20260101-000000-abc123",
        original_text=render_markdown(original), structured=original,
        report=run_ats_checks(render_markdown(original), None, original),
        tailored=TailoredResume(
            resume=StructuredResume.model_validate(unfinished), changes=[], warnings=[]),
    )
    server._save_resume_doc(doc)
    return TestClient(server.app), doc.resume_id


def _get(client, rid, **params):
    return client.get(f"/resumes/{rid}/download", params=params)


@pytest.mark.parametrize("fmt", FORMATS)
def test_a_tailored_export_with_placeholders_is_refused_in_every_format(desk, fmt):
    client, rid = desk

    r = _get(client, rid, fmt=fmt, version="tailored")

    assert r.status_code == 409, f"{fmt} was served with a placeholder in it"
    detail = r.json()["detail"]
    assert "[METRIC]" in detail and "Initrode" in detail, \
        "the refusal must say what is unresolved and where"


@pytest.mark.parametrize("fmt", FORMATS)
def test_the_original_stays_downloadable(desk, fmt):
    client, rid = desk

    r = _get(client, rid, fmt=fmt, version="original")

    assert r.status_code == 200
    assert r.headers["x-scrivio-export"] == "original"


def test_filling_the_number_opens_the_export_and_it_contains_that_number(desk):
    client, rid = desk
    filled = client.post(f"/resumes/{rid}/fill-metrics", json={"values": ["30%"]})
    assert filled.status_code == 200

    md = _get(client, rid, fmt="md", version="tailored")
    word = _get(client, rid, fmt="docx", version="tailored")

    assert md.status_code == 200 and word.status_code == 200
    assert md.headers["x-scrivio-export"] == "final"
    assert "reducing costs by 30%" in md.text and "[METRIC]" not in md.text
    paragraphs = [p.text for p in docx.Document(io.BytesIO(word.content)).paragraphs]
    assert any("reducing costs by 30%" in p for p in paragraphs), \
        "every format renders the same saved content"


def test_a_draft_is_available_only_when_asked_for_and_is_labelled_as_one(desk):
    client, rid = desk

    r = _get(client, rid, fmt="md", version="tailored", draft="true")

    assert r.status_code == 200
    assert r.headers["x-scrivio-export"] == "draft"
    assert "DRAFT" in r.headers["content-disposition"]
    assert "[METRIC]" in r.text


def test_an_export_of_a_page_that_has_gone_stale_is_refused(desk):
    """The server cannot see unsaved text in a browser, but it can refuse
    to hand over a version other than the one the page is showing."""
    client, rid = desk
    shown = client.get(f"/resumes/{rid}").json()["updated_at"]
    client.post(f"/resumes/{rid}/fill-metrics", json={"values": ["30%"]})

    stale = _get(client, rid, fmt="md", version="tailored", expect=shown)
    fresh = _get(client, rid, fmt="md", version="tailored",
                 expect=client.get(f"/resumes/{rid}").json()["updated_at"])

    assert stale.status_code == 409 and "changed" in stale.json()["detail"]
    assert fresh.status_code == 200


def test_an_export_is_refused_while_tailoring_is_still_running(desk, monkeypatch):
    client, rid = desk
    # Running in this process. Without this the status alone would be read
    # as left over from a server that stopped.
    monkeypatch.setattr(server, "_RESUME_WORK", {rid})
    doc = server._load_resume_doc(rid)
    doc.tailor_status = "tailoring"
    server._save_resume_doc(doc)

    r = _get(client, rid, fmt="md", version="tailored", draft="true")

    assert r.status_code == 409


def test_a_placeholder_outside_the_bullets_blocks_too(desk):
    """Headline and user-made sections were not in the old placeholder
    scan, so a placeholder there was invisible to every gate."""
    client, rid = desk
    client.post(f"/resumes/{rid}/fill-metrics", json={"values": ["30%"]})
    doc = server._load_resume_doc(rid)
    doc.tailored.resume.basics.label = "Engineer with [METRIC] years of experience"
    server._save_resume_doc(doc)

    r = _get(client, rid, fmt="md", version="tailored")

    assert r.status_code == 409 and "headline" in r.json()["detail"]
